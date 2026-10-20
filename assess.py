"""SonarQube maturity assessment with Jev (TypeSafe System One), for any customer.

Each customer has a folder under customers/ (create one with --new):
  customers/<name>/
    customer.json        optional: name, confirmed facts, extra bullets, overrides
    *notes*.txt|md|docx  assessor notes (any file with "note" in its name)
    *.txt|md|vtt|docx    everything else is treated as meeting transcript(s)
    metrics/*.csv        telemetry exports (Languages, Projects, LoC Usage, Users, ...)
    *.json               optional SonarQube system info / support info export

Jev judges, from that evidence:
  * which maturity criteria hold (Nouls) -> code derives the level,
  * a holistic maturity Score as a cross-check,
  * whether each recommendation applies and is already in place,
  * which executive-summary, environment and Quality Gate statements are supported,
  * the SonarQube version, edition, database and login method (Choices).
Code then fills in the Maturity Assessment Template docx.

Usage:
  .venv/bin/python assess.py --new "Acme Corp" --assessor "Jane Doe"   # create customers/acme_corp/
  .venv/bin/python assess.py acme                   # call Jev, build docx + report
  .venv/bin/python assess.py acme --from-cache      # rebuild without calling Jev
  .venv/bin/python assess.py --list                 # list customer folders
  .venv/bin/streamlit run app.py                    # the same, in a web page
"""

import argparse
import copy
import json
import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

import library
import serverinfo
import telemetry

HERE = Path(__file__).parent
TEMPLATES = HERE / "templates"
TEMPLATE = TEMPLATES / "Maturity Assessment Template.docx"
RECS_MD = TEMPLATES / "Maturity Assessment Recommendations by Level.md"
CUSTOMERS = HERE / "customers"
OUTPUT_DIR = HERE / "reports"  # where the finished docx is written
DOC_SUFFIXES = {".txt", ".md", ".vtt", ".docx"}
# Jev's default limits (docs.typesafe.ai/models): 32k tokens for the state plus any one
# question, 64k for the state plus all questions. Budgets leave a margin under each.
TOKEN_BUDGET = 30_000    # state + the longest question (limit 32k)
REQUEST_BUDGET = 60_000  # state + every question (limit 64k)

YES = 0.5
REVIEW_BAND = (0.3, 0.7)  # Nouls in this band are flagged for a human look


class AssessError(Exception):
    """A problem with the inputs or setup, worded for the person running the tool."""


def load_env():
    """Fill os.environ from .env (TYPESAFE_API_KEY, MATURITY_ASSESSOR) without overriding it."""
    env = HERE / ".env"
    if env.exists():
        for line in env.read_text().splitlines():
            if "=" in line and not line.lstrip().startswith("#"):
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip("'\""))

LEVELS = {
    1: ("External Audits", "Individuals outside of the development process are using SonarQube to do one-off audits. Results may be used to create a backlog of items to fix at a later date."),
    2: ("Spot Checks Within Teams", "Sonar tools are in use within individual teams. Developers begin to have the opportunity to act on issues and see the effects of their code changes."),
    3: ("Automated Code Checks", "Analysis of code is automatically triggered via integration with a continuous integration (CI) process. Developers are positioned to fix issues as they are created."),
    4: ("Standards Enforcement", "Quality standards are enforced within the development pipeline. Code is not allowed to be promoted if it does not meet standards. Organizations begin to decrease the accumulation of technical debt."),
    5: ("Standards Governance", "Quality standards are enforced and governed centrally across the organization. Technical debt is reduced to manageable levels and Clean as You Code is fully implemented at scale. Leadership has access to up-to-date reporting and can make data-driven decisions to improve code quality."),
}

def evidence(state):
    fields = [f"`{k}`" for k in ("assessor_confirmed", "assessor_notes", "meeting_transcript", "server_info", "telemetry") if k in state]
    text = f"Use only the evidence in {', '.join(fields)}."
    if "assessor_confirmed" in state:
        text += " `assessor_confirmed` lists facts the Sonar assessor verified and overrides anything else."
    if "server_info" in state:
        text += " `server_info` comes from the customer's SonarQube system information export and is authoritative for how the instance is configured today."
    if "meeting_transcript" in state:
        text += (f" In the transcript, the Sonar assessor appears as \"You\" or as {state['assessor']}; "
                 "all other speakers are the customer or the Sonar account team.")
    return text

# Maturity criteria: level -> [(id, condition, true, false)]. Level N is reached
# when every criterion of levels 2..N holds.
LEVEL_CRITERIA = {
    2: [("teams_use_results",
         "Do the customer's development teams themselves use SonarQube analysis results in their work (not only outside auditors doing one-off audits)?",
         "Developers or teams look at and act on SonarQube results.",
         "Only people outside development run occasional audits, or there is no evidence of team usage.")],
    3: [("ci_automated",
         "Is SonarQube analysis triggered automatically by CI pipelines for the customer's projects?",
         "Analysis runs automatically as part of CI for most projects.",
         "Analysis is run manually, ad hoc, or there is no evidence of CI automation."),
        ("pr_analysis",
         "Are pull requests analyzed by SonarQube with results visible in the customer's DevOps platform?",
         "Pull request analysis is in regular use and visible in the DevOps platform.",
         "Pull requests are not analyzed, or there is no evidence they are.")],
    4: [("qg_enforced",
         "Does a failing SonarQube Quality Gate actually block merges or fail pipelines for the customer's projects?",
         "A failing Quality Gate stops code from being merged or promoted for most projects.",
         "The Quality Gate is informational, a soft failure, optional, or only enforced by a few teams.")],
    5: [("central_governance",
         "Are Quality Gates and Quality Profiles governed centrally, with all teams using shared organization standards instead of their own custom gates and profiles?",
         "One central function owns the gates and profiles and teams use the shared standards.",
         "Some teams create or maintain their own custom gates or profiles."),
        ("leadership_portfolios",
         "Do the customer's leaders regularly use SonarQube portfolios with up-to-date metrics to make decisions?",
         "Leadership actively reviews portfolio reporting for decisions.",
         "Portfolios are unused, lightly used, or leadership does not yet get this reporting.")],
}

# Recommendations that exist in the template but not in the recommendations md.
TEMPLATE_ONLY_RECS = [
    {"key": "rationalize_qg", "level": 5, "section": None,
     "title": "Rationalize Quality Gates for proper governance",
     "description": "Define a base default Quality Gate focused on new code that is enforced for all projects, plus 2-4 incrementally stricter gates, and remove inconsistent or project-specific gate conditions."},
    {"key": "portfolio_reporting", "level": 5, "section": None,
     "title": "Use Portfolio reporting as part of regular governance reviews",
     "description": "Use Portfolios to give leadership, security, compliance and audit stakeholders a consolidated view of releasability, security, reliability and maintainability across projects, including periodic PDF and security reports."},
]

# Template "quality gate inconsistencies" bullets -> paragraph prefixes they own.
QG_FINDINGS = {
    "zero_coverage": ("Some of the customer's Quality Gates use a 0% coverage or 0% security hotspots reviewed condition.", ["0% coverage as a criteria"]),
    "looser_security": ("Some of the customer's Quality Gates require a looser rating on security than on reliability or maintainability.", ["Looser criteria on security"]),
    "loose_new_code": ("Some of the customer's Quality Gates use very loose conditions on new code, such as a maintainability or reliability rating of D.", ["Too loose criteria on new code"]),
    "absolute_counts": ("Some of the customer's Quality Gates use absolute-count conditions such as duplicated blocks or duplicated files.", ["Inconsistency criteria", "Duplicated blocks:", "Duplicated files:"]),
}

# Recommendation (md section title or rec key) -> maturity criteria it exists to
# achieve. When Jev scores all of them as already met, the recommendation is dropped.
REC_CRITERIA = {
    "Integrate with DevOps Platform(s)": ["ci_automated", "pr_analysis"],
    "Automate Code Scanning": ["ci_automated", "pr_analysis"],
    "Enforce Quality Gates": ["qg_enforced"],
    "Consolidate Quality Standards": ["central_governance"],
    "rationalize_qg": ["central_governance"],
    "portfolio_reporting": ["leadership_portfolios"],
}

# Recommendation key -> environment statement that means it is already done. The
# generic "done" question asks about all teams and projects, which does not fit
# instance-wide setup such as a DevOps platform integration.
REC_DONE_IF_ENV = {
    "integrate_with_github": "env_installation_github",
    "integrate_with_azure_devops": "env_installation_azure_devops",
}

TEMPLATE_HIGHLIGHTS = {"ff9900"}  # template placeholder shading, removed from the output

# Template recommendation markers (exact paragraph text) in document order.
TEMPLATE_SECTIONS = {
    "Integrate with DevOps Platform(s)": ["Integrate with GitHub", "Integrate with Azure DevOps"],
    "Automate Code Scanning": ["Automate Azure Pipeline branch scanning", "Publish Quality Gate results"],
    "Rationalize Quality Gates for proper governance": [],  # whole section is one recommendation
    "Validate New Code Settings": ["Set default New Code setting", "Educate Project Administrators on New Code options"],
    "Enforce Quality Gates": ["Add Branch Protection Rules to GitHub repositories", "Stop code promotion when Quality Gate is failing"],
    "Stimulate the adoption of SonarQube for IDE Usage": ["Promote SonarQube for IDE adoption", "Track SonarQube for IDE usage"],
    "Reporting Management": ["Use Portfolio reporting as part of regular governance reviews"],
}


def norm(s):
    return re.sub(r"\s+", " ", s.replace("“", '"').replace("”", '"').replace("’", "'")).strip().lower()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", norm(s)).strip("_")[:48]


def strip_md_links(s):
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s).replace("**", "")


# ---------------------------------------------------------------- inputs

def parse_recommendations_md(text):
    """'# To Achieve Level N' > '### Section' (intro text) > '- [ ] Item' (indented description)."""
    recs, sections = [], {}
    level, section, item = None, None, None
    for raw in text.splitlines():
        line = raw.strip()
        if m := re.match(r"^# To Achieve Level (\d)", line):
            level, section, item = int(m.group(1)), None, None
        elif level is None:
            continue
        elif line.startswith("### ") and not raw.startswith(" "):
            section, item = line[4:].strip(), None
            sections[section] = {"level": level, "intro": ""}
        elif line.startswith("- [ ]"):
            title = re.sub(r"^#+\s*", "", line[5:].strip())
            item = {"key": slug(title), "level": level, "section": section, "title": strip_md_links(title), "description": ""}
            recs.append(item)
        elif line:
            line = re.sub(r"^#+\s*", "", line)
            if item is not None:
                item["description"] = (item["description"] + " " + line).strip()
            elif section is not None:
                sections[section]["intro"] = (sections[section]["intro"] + " " + line).strip()
    return recs, sections


def read_doc(path):
    if path.suffix == ".docx":
        return "\n".join(p.text for p in docx.Document(path).paragraphs)
    return path.read_text(errors="replace")


CUE_TIMING = re.compile(r"^(?:\d+:)?\d{1,2}:\d{2}(?:[.,]\d+)?\s*-->\s*(?:\d+:)?\d{1,2}:\d{2}(?:[.,]\d+)?\b.*$")
VOICE_TAG = re.compile(r"^<v(?:\.[^ >]*)?\s+([^>]+)>(.*?)(?:</v>)?$")
SPEAKER = re.compile(r"^([^:]{1,60}?):\s+(.*)$")


def clean_transcript(text):
    """Caption-style transcript (.vtt, or .txt with cue timings) -> one line per speaker turn.

    Drops cue timings, cue numbers, WEBVTT/NOTE headers and blank lines, and merges
    consecutive captions from the same speaker. The spoken words are kept as they are.
    Timings cost Jev many tokens and carry nothing it judges. Other text is returned unchanged.
    """
    lines = [l.strip() for l in text.splitlines()]
    if sum(bool(CUE_TIMING.match(l)) for l in lines) < 3:
        return text
    captions, in_note = [], False
    for i, l in enumerate(lines):
        if not l:
            in_note = False
        elif in_note or l.split(" ", 1)[0] in ("NOTE", "STYLE", "REGION"):
            in_note = True  # VTT comment/style block, ends at the next blank line
        elif l == "WEBVTT" or l.startswith("WEBVTT ") or CUE_TIMING.match(l):
            continue
        elif l.isdigit() and i + 1 < len(lines) and CUE_TIMING.match(lines[i + 1]):
            continue  # cue number
        elif m := VOICE_TAG.match(l):
            captions.append((m[1].strip(), strip_tags(m[2]).strip()))
        else:
            captions.append((None, strip_tags(l)))
    # A "Name:" prefix counts as a speaker only when it recurs, so "So the answer: yes" stays text.
    prefixes = [m[1] for who, c in captions if who is None and (m := SPEAKER.match(c))]
    named = {s for s in prefixes if prefixes.count(s) >= 3}
    turns = []  # [speaker, [what they said]]; speaker None = a caption kept on its own line
    for who, said in captions:
        if who is None and (m := SPEAKER.match(said)) and m[1] in named:
            who, said = m[1], m[2]
        if who is not None and turns and turns[-1][0] == who:
            turns[-1][1].append(said)  # only the same speaker's consecutive captions are merged
        else:
            turns.append([who, [said]])
    return "\n".join(f"{who}: {' '.join(said)}" if who else said[0] for who, said in turns)


def strip_tags(caption):
    """VTT cue markup (<c.yellow>, <i>, <00:00:01.000>) around the spoken words."""
    return re.sub(r"</?(?:c|i|b|u|v|lang|ruby|rt)\b[^>]*>|<\d[\d:.]*>", "", caption)


def find_customer(arg):
    path = Path(arg).expanduser()
    if path.is_dir():
        return path
    folders = sorted(d for d in CUSTOMERS.glob("*") if d.is_dir()) if CUSTOMERS.is_dir() else []
    exact = [d for d in folders if d.name == slug(arg)]
    partial = [d for d in folders if slug(arg) in d.name]
    if exact or len(partial) == 1:
        return (exact or partial)[0]
    names = ", ".join(d.name for d in folders) or "none yet"
    raise AssessError(f"No unique customer folder matches {arg!r} (customers: {names}). Create one with --new.")


def list_customers():
    return sorted(d.name for d in CUSTOMERS.glob("*") if d.is_dir()) if CUSTOMERS.is_dir() else []


def new_customer(name, assessor=None):
    folder = CUSTOMERS / slug(name)
    (folder / "metrics").mkdir(parents=True, exist_ok=True)
    cfg = folder / "customer.json"
    if not cfg.exists():
        load_env()
        cfg.write_text(json.dumps({
            "customer": name,
            "assessor": assessor or os.environ.get("MATURITY_ASSESSOR", ""),
            "confirmed": [],
            "environment_extra": {},
            "summary_extra": {},
            "exclude_recs": [],
            "include_recs": [],
        }, indent=2) + "\n")
    print(f"""Created {folder}
Next:
  1. Put your notes in a file with "notes" in its name (e.g. notes.txt), and the meeting transcript(s) next to it.
  2. Copy the telemetry CSV exports into {folder / "metrics"}.
  3. Check "assessor" in customer.json (your name, as it appears in the transcript) and
     optionally list facts you verified under "confirmed".
  4. Run: .venv/bin/python assess.py {folder.name}""")
    return folder


def read_config(folder):
    cfg_path = folder / "customer.json"
    try:
        return json.loads(cfg_path.read_text()) if cfg_path.exists() else {}
    except json.JSONDecodeError as e:
        raise AssessError(f"{cfg_path} is not valid JSON: {e}")


def customer_name(folder, cfg=None):
    cfg = read_config(folder) if cfg is None else cfg
    return cfg.get("customer") or folder.name.replace("_", " ").title()


def load_customer(folder):
    """Everything the run needs for one customer, discovered from its folder."""
    load_env()
    cfg_path = folder / "customer.json"
    cfg = read_config(folder)
    name = customer_name(folder, cfg)
    assessor = cfg.get("assessor") or os.environ.get("MATURITY_ASSESSOR")
    if not assessor:
        raise AssessError(f'No assessor name: set "assessor" in {cfg_path} or MATURITY_ASSESSOR in .env.')
    docs = sorted(f for f in folder.iterdir() if f.is_file() and f.suffix.lower() in DOC_SUFFIXES and not f.name.startswith((".", "~$")))
    notes = [f for f in docs if "note" in f.name.lower()]
    transcripts = [f for f in docs if f not in notes]
    metrics = folder / "metrics" if (folder / "metrics").is_dir() else folder
    tel, tel_bullets = telemetry.load(metrics)
    tel.update(cfg.get("telemetry_extra", {}))
    info_file, server_info, server_choices = serverinfo.load(folder)
    out = folder / "output"
    out.mkdir(exist_ok=True)
    return {
        "name": name,
        "folder": folder,
        "assessor": assessor,
        "confirmed": cfg.get("confirmed", []),
        "notes": "\n\n".join(f"### {f.name}\n{read_doc(f)}" if len(notes) > 1 else read_doc(f) for f in notes),
        "transcript": "\n\n".join(f"### {f.name}\n{clean_transcript(read_doc(f))}" if len(transcripts) > 1
                                  else clean_transcript(read_doc(f)) for f in transcripts),
        "files": {"notes": [f.name for f in notes], "transcripts": [f.name for f in transcripts],
                  "metrics": sorted(f.name for f in metrics.glob("*.csv")), "server_info": info_file},
        "server_info": server_info,
        "server_choices": server_choices,
        "telemetry": tel,
        "telemetry_bullets": tel_bullets,
        "environment": cfg.get("environment"),
        "environment_extra": cfg.get("environment_extra", {}),
        "summary_extra": cfg.get("summary_extra", {}),
        "exclude_recs": set(cfg.get("exclude_recs", [])),
        "include_recs": set(cfg.get("include_recs", [])),
        "out_docx": OUTPUT_DIR / f"{name} - Maturity Assessment.docx",
        "out_report": out / f"{slug(name)}_maturity_results.md",
        "cache": out / "jev_answers.json",
    }


def build_state(ctx):
    state = {"customer": ctx["name"], "assessor": ctx["assessor"]}
    if ctx["confirmed"]:
        state["assessor_confirmed"] = ctx["confirmed"]
    if ctx["notes"].strip():
        state["assessor_notes"] = ctx["notes"]
    if ctx["transcript"].strip():
        state["meeting_transcript"] = ctx["transcript"]
    if ctx["server_info"]:
        state["server_info"] = ctx["server_info"]
    if ctx["telemetry"]:
        state["telemetry"] = ctx["telemetry"]
    if len(state) == 2:
        raise AssessError(f"No notes, transcript, confirmed facts or metrics found in {ctx['folder']}.")
    return state


VERSION_RE = re.compile(r"\b(20[2-9]\d\.\d{1,2}(?:\.\d{1,2})?|(?:8|9|10)\.\d{1,2}(?:\.\d{1,2})?)\b")


def version_candidates(state):
    """SonarQube version strings mentioned anywhere in the evidence (Jev picks one)."""
    def strings(v):
        if isinstance(v, str):
            yield v
        elif isinstance(v, (list, tuple)):
            for x in v:
                yield from strings(x)
        elif isinstance(v, dict):
            for x in v.values():
                yield from strings(x)
    text = "\n".join(strings({k: v for k, v in state.items() if k != "telemetry"}))
    return sorted(set(VERSION_RE.findall(text)))


def fmt(statement, ctx, version=None):
    text = statement.replace("{C}", ctx["name"]).replace("{latest}", library.LATEST_LTA)
    if version:
        return text.replace("{version}", version)
    return re.sub(r"\s*\(\{version\}\)", "", text).replace("{version}", "its current version")


def summary_candidates(ctx):
    return {**library.SUMMARY_CANDIDATES, **ctx["summary_extra"]}


# ---------------------------------------------------------------- questions

def build_questions(recs, state, ctx):
    from typesafe_sdk import Choice, Noul, NoulCriteria, Score

    EVIDENCE = evidence(state)
    qs = {}
    for level, crits in LEVEL_CRITERIA.items():
        for cid, cond, t, f in crits:
            qs[f"crit_{cid}"] = Noul(instructions=f"{cond} {EVIDENCE}", criteria=NoulCriteria(true=t, false=f))

    qs["maturity_level"] = Score(
        instructions=f"Which Sonar maturity level best describes how the customer uses SonarQube today? Pick the highest level whose description is fully true for the organization as a whole. {EVIDENCE}",
        criteria=[f"Level {n} - {name}: {desc}" for n, (name, desc) in LEVELS.items()],
    )

    for r in recs:
        rec = {"title": r["title"], "description": strip_md_links(fmt(r["description"], ctx))}
        qs[f"rec_{r['key']}_applies"] = Noul(
            instructions={"recommendation": rec,
                          "question": f"Does this recommendation apply to the customer's environment (their DevOps platform, tools and way of working)? {EVIDENCE}"},
            criteria=NoulCriteria(true="The recommendation targets tools or practices the customer has or would adopt.",
                                  false="The recommendation targets a platform or tool the customer does not use."))
        done = r.get("done") or {
            "question": "Has the customer already fully implemented this recommendation across the organization?",
            "true": "The evidence shows the recommended practice is already in place for essentially all teams and projects.",
            "false": "The practice is missing, partial, optional, used by only some teams, or there is no evidence it is in place."}
        qs[f"rec_{r['key']}_done"] = Noul(
            instructions={"recommendation": rec, "question": f"{done['question']} {EVIDENCE}"},
            criteria=NoulCriteria(true=done["true"], false=done["false"]))

    for key, statement in summary_candidates(ctx).items():
        qs[f"sum_{key}"] = Noul(
            instructions={"statement": fmt(statement, ctx),
                          "question": f"Is this statement about the customer supported by the evidence? {EVIDENCE}"},
            criteria=NoulCriteria(true="The evidence directly supports the statement.",
                                  false="The evidence contradicts the statement or does not mention it."))

    for key, (finding, _) in QG_FINDINGS.items():
        qs[f"qg_{key}"] = Noul(
            instructions={"finding": finding, "question": f"Does the evidence show this finding is true? {EVIDENCE}"},
            criteria=NoulCriteria(true="The evidence explicitly describes this gate condition.",
                                  false="The evidence does not describe this gate condition."))

    if not ctx["environment"]:  # environment section is generated, not supplied
        for heading, items in library.ENV_CANDIDATES.items():
            for key, statement in items.items():
                qs[f"env_{slug(heading)}_{key}"] = Noul(
                    instructions={"statement": statement, "question": f"Is this statement about the customer's current SonarQube setup supported by the evidence? {EVIDENCE}"},
                    criteria=NoulCriteria(true="The evidence directly supports the statement for the customer today.",
                                          false="The evidence contradicts the statement, describes only a plan or option, or does not mention it."))
        for key, (question, options) in library.ENV_CHOICES.items():
            if key in ctx["server_choices"]:
                continue  # taken from the server info export
            qs[f"choice_{key}"] = Choice(instructions=f"{question} {EVIDENCE}", criteria=options)
    if "version" not in ctx["server_choices"] and (versions := version_candidates(state)):
        qs["choice_version"] = Choice(
            instructions=f"Which SonarQube Server version is the customer running today? Versions mentioned only as upgrade targets, new releases or examples do not count. {EVIDENCE}",
            criteria={**{v: f"The customer currently runs SonarQube {v}" for v in versions},
                      "unknown": "The evidence does not say which version the customer runs today"})
    return qs


def ask_jev(state, questions):
    from typesafe_sdk import TypeSafeClient

    load_env()
    if not os.environ.get("TYPESAFE_API_KEY"):
        raise AssessError(f"TYPESAFE_API_KEY is not set: add it to {HERE / '.env'} (see .env.bk).")
    with TypeSafeClient() as client:
        resp = client.system_one(state=state, questions=questions)
    answers = {}
    for qid, a in resp.answers.items():
        d = a.model_dump() if hasattr(a, "model_dump") else dict(a)
        answers[qid] = {k: v for k, v in d.items() if k in ("type", "noul", "choice", "score", "confidence", "probabilities")}
    meta = {"model": resp.model, "request_id": getattr(resp, "request_id", None), "usage": resp.usage.model_dump() if hasattr(resp.usage, "model_dump") else str(resp.usage)}
    return answers, meta


# ---------------------------------------------------------------- decisions

def decide(answers, recs, ctx):
    p = lambda qid: answers[qid]["noul"]
    losers = set()
    for group in library.EXCLUSIVE:
        present = [q for q in group if q in answers]
        if present:
            losers |= set(present) - {max(present, key=p)}
    holds = lambda qid: qid in answers and qid not in losers and p(qid) > YES

    level = 1
    for n in sorted(LEVEL_CRITERIA):
        if all(p(f"crit_{cid}") > YES for cid, *_ in LEVEL_CRITERIA[n]):
            level = n
        else:
            break

    score = answers["maturity_level"]
    score_level = max(score["probabilities"], key=lambda k: score["probabilities"][k])
    score_level = int(score_level) + 1  # criteria index 0 == Level 1

    chosen = []
    for r in recs:
        applies, done = p(f"rec_{r['key']}_applies"), p(f"rec_{r['key']}_done")
        if (env := REC_DONE_IF_ENV.get(r["key"])) and env in answers:
            done = max(done, p(env))
        crits = REC_CRITERIA.get(r["key"]) or REC_CRITERIA.get(r["section"]) or []
        met = [c for c in crits if p(f"crit_{c}") > YES]
        achieved = bool(crits) and len(met) == len(crits)
        if r["key"] in ctx["exclude_recs"]:
            reason = "excluded in customer.json"
        elif r["key"] in ctx["include_recs"]:
            reason = "forced by customer.json"
        elif achieved:
            reason = "criteria already met: " + ", ".join(met)
        elif applies <= YES:
            reason = "does not apply"
        elif done > YES:
            reason = "already done"
        else:
            reason = ""
        r = dict(r, p_applies=applies, p_done=done, criteria=crits, reason=reason,
                 include=r["key"] in ctx["include_recs"] or (
                     r["key"] not in ctx["exclude_recs"] and applies > YES and done <= YES and not achieved))
        r["review"] = any(REVIEW_BAND[0] < x < REVIEW_BAND[1] for x in (applies, done))
        chosen.append(r)

    choices = {k[len("choice_"):]: v["choice"] for k, v in answers.items() if k.startswith("choice_") and v.get("choice") != "unknown"}
    choices.update(ctx["server_choices"])
    summary = [k for k in summary_candidates(ctx) if holds(f"sum_{k}")]
    if "version" not in choices and "current_version" in summary:
        summary.remove("current_version")
    qg = [k for k in QG_FINDINGS if p(f"qg_{k}") > YES]
    return {"level": level, "score_level": score_level, "score": score, "recs": chosen, "summary": summary, "qg": qg,
            "choices": choices, "environment": environment(answers, holds, choices, ctx)}


def environment(answers, holds, choices, ctx):
    """'<Customer> Current Environment' bullets: {heading: [(text, indent)]}."""
    if ctx["environment"]:
        return {h: [tuple(b) for b in items] for h, items in ctx["environment"].items()}
    env = {h: [] for h in ("Installation", "Utilization", "Code Analysis", "Governance")}
    product = " ".join(x for x in ("SonarQube Server", choices.get("edition"), choices.get("version")) if x)
    if "database" in choices:
        product += f", {choices['database']} database"
    env["Installation"].append((product, 0))
    if "authentication" in choices:
        env["Installation"].append((f"Users authenticate through {choices['authentication']}", 0))
    for heading, items in library.ENV_CANDIDATES.items():
        env[heading] += [(text, 0) for key, text in items.items() if holds(f"env_{slug(heading)}_{key}")]
    for heading, items in ctx["telemetry_bullets"].items():
        env[heading] += items
    for heading, items in ctx["environment_extra"].items():
        env.setdefault(heading, []).extend(tuple(b) for b in items)
    return {h: items for h, items in env.items() if items}


# ---------------------------------------------------------------- docx helpers

def set_text(p_el, text):
    """Replace a paragraph's text, keeping the formatting of its first run."""
    runs = p_el.findall(".//" + qn("w:r"))
    first = copy.deepcopy(runs[0]) if runs else OxmlElement("w:r")
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)
    for child in list(first):
        if child.tag != qn("w:rPr"):
            first.remove(child)
    p_el.append(first)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    first.append(t)
    return p_el


def set_text_with_links(p_el, text, part):
    """Like set_text, but turns markdown [label](url) into real hyperlinks."""
    set_text(p_el, "")
    template_run = p_el.findall(qn("w:r"))[-1]
    p_el.remove(template_run)
    for i, chunk in enumerate(re.split(r"(\[[^\]]+\]\([^)]+\))", text.replace("**", ""))):
        if not chunk:
            continue
        run = copy.deepcopy(template_run)
        t = run.find(qn("w:t"))
        if m := re.match(r"\[([^\]]+)\]\(([^)]+)\)", chunk):
            t.text = m.group(1)
            rpr = run.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                run.insert(0, rpr)
            style = OxmlElement("w:rStyle")
            style.set(qn("w:val"), "Hyperlink")
            rpr.insert(0, style)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), part.relate_to(m.group(2), RT.HYPERLINK, is_external=True))
            link.append(run)
            p_el.append(link)
        else:
            t.text = chunk
            p_el.append(run)


def ptext(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def replace_placeholders(doc, mapping):
    def paragraphs():
        yield from doc.paragraphs
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        for s in doc.sections:
            for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                yield from hf.paragraphs

    for para in paragraphs():
        if not any(k in para.text for k in mapping):
            continue
        for run in para.runs:
            for k, v in mapping.items():
                run.text = run.text.replace(k, v)
        if any(k in para.text for k in mapping):  # placeholder split across runs
            text = para.text
            for k, v in mapping.items():
                text = text.replace(k, v)
            set_text(para._p, text)


WPS = "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}"
DML = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def fit_level_title(el):
    """The level graphic's title is a fixed-size text box (no autofit) with trailing
    empty lines; with a fallback font the name wraps and spills over the graphic.
    Drop the empty lines and size the font so the longest line fits on one line."""
    for wsp in el.iter(WPS + "wsp"):
        box = wsp.find(WPS + "txbx")
        ext = wsp.find(f"{WPS}spPr/{DML}xfrm/{DML}ext")
        if box is None or ext is None or not re.match(r"Level \d:", ptext(box)):
            continue
        paras = list(box.iter(qn("w:p")))
        for para in paras:
            if not ptext(para).strip():
                para.getparent().remove(para)
        lines = [ptext(para) for para in paras if ptext(para).strip()]
        width_pt = int(ext.get("cx")) / 12700 - 2 * 7.2  # minus the box's 0.1" side insets
        fit = int(2 * width_pt / (max(map(len, lines)) * 0.68))  # half-points; ~0.68 em per char for a semibold sans
        for sz in box.iter(qn("w:sz"), qn("w:szCs")):
            sz.set(qn("w:val"), str(min(int(sz.get(qn("w:val"))), fit)))


def build_docx(result, recs_md_sections, ctx):
    doc = docx.Document(TEMPLATE)
    body = doc.element.body
    C = ctx["name"]
    version = result["choices"].get("version")
    recs_md_sections = {**recs_md_sections, **{s["section"]: {"intro": s["intro"]} for s in library.EXTRA_SECTIONS}}

    replace_placeholders(doc, {"<SONARSOURCER>": ctx["assessor"], "<CUSTOMER>": C, "<Customer>": C})

    def children():
        return list(body.iterchildren())

    def find(text, start=None, exact=True):
        kids = children()
        i0 = kids.index(start) if start is not None else 0
        for el in kids[i0:]:
            if el.tag == qn("w:p") and ((ptext(el) == text) if exact else ptext(el).startswith(text)):
                return el
        raise KeyError(text)

    def between(a, b):
        kids = children()
        return kids[kids.index(a) + 1: kids.index(b)]

    level, (lname, _) = result["level"], LEVELS[result["level"]]

    # Executive summary: replace guidance + example bullets with selected statements.
    intro = find("Sonar and ", exact=False)
    bullets = [el for el in between(intro, find("The Sonar Maturity Model")) if el.find(".//" + qn("w:numPr")) is not None]
    sample = bullets[0]
    lines = [f"{C} is currently at Level {level}: {lname}."]
    if level < 5:
        lines[0] += f" The recommended next step is Level {level + 1}: {LEVELS[level + 1][0]}."
    lines += [fmt(summary_candidates(ctx)[k], ctx, version) for k in result["summary"]]
    for line in lines:
        sample.addprevious(set_text(copy.deepcopy(sample), line))
    for el in bullets:
        body.remove(el)

    # Current environment.
    env_h1 = find(f"{C} Current Environment")
    assess_h1 = find("Assessment")
    old = between(env_h1, assess_h1)
    h2 = next(el for el in old if el.tag == qn("w:p") and ptext(el) == "Installation")
    lvl0 = next(el for el in old if ptext(el).startswith("SonarQube Enterprise Edition"))
    lvl1 = next(el for el in old if ptext(el).startswith("License entitlement"))
    for heading, items in result["environment"].items():
        assess_h1.addprevious(set_text(copy.deepcopy(h2), heading))
        for text, indent in items:
            assess_h1.addprevious(set_text(copy.deepcopy(lvl1 if indent else lvl0), text))
    for el in old:
        body.remove(el)

    # Current maturity level: keep only the matching level block.
    rec_h1 = find("Recommendations")
    blocks, current = {}, None
    for el in between(find("Current Maturity Level"), rec_h1):
        if el.tag == qn("w:p") and (m := re.match(r"^Level (\d): ", ptext(el))):
            current = int(m.group(1))
        if current:
            blocks.setdefault(current, []).append(el)
    keep = level if level in blocks else min(blocks)
    if keep != level:  # Level 1 has no benefits table in the template
        set_text(blocks[keep][0], f"Level {level}: {lname}")
        for el in blocks[keep][1:]:
            body.remove(el)
    for n, els in blocks.items():
        if n != keep:
            for el in els:
                body.remove(el)
    for el in blocks[keep]:
        fit_level_title(el)

    # Recommendation intro sentence.
    rec_intro = find("The next step in your Sonar evolution", exact=False)
    if level < 5:
        set_text(rec_intro, f"The next step in your Sonar evolution would be {LEVELS[level + 1][0]}. Here are recommendations for how to get there:")
    else:
        set_text(rec_intro, "Here are recommendations to sustain Standards Governance at scale:")

    # Quality Gate inconsistencies: drop findings without evidence.
    for key, (_, prefixes) in QG_FINDINGS.items():
        if key not in result["qg"]:
            for pre in prefixes:
                body.remove(find(pre, start=rec_h1, exact=False))
    if not result["qg"]:
        body.remove(find("In the case of ", start=rec_h1, exact=False))

    # Recommendation units: [marker, next marker).
    annex = find("Annexes")
    by_title = {norm(r["title"]): r for r in result["recs"]}
    markers = []
    for sec, items in TEMPLATE_SECTIONS.items():
        markers.append(("section", sec, find(sec, start=rec_h1)))
        markers += [("item", it, find(it, start=rec_h1)) for it in items]
    kids = children()
    markers.sort(key=lambda m: kids.index(m[2]))
    units = {}
    for i, (kind, title, el) in enumerate(markers):
        end = markers[i + 1][2] if i + 1 < len(markers) else annex
        units[title] = (kind, between_inclusive(kids, el, end))

    removed_items = set()
    for sec, items in TEMPLATE_SECTIONS.items():
        titles = items or [sec]
        for it in titles:
            r = by_title[norm(it)]
            if not r["include"]:
                removed_items.add(it)
        if all(it in removed_items for it in titles):
            for el in units[sec][1]:
                body.remove(el)
        for it in items:
            if it in removed_items:
                for el in units[it][1]:
                    body.remove(el)

    # Recommendations from the md that the template does not contain.
    template_titles = {norm(t) for sec, its in TEMPLATE_SECTIONS.items() for t in (its or [sec])}
    extra = [r for r in result["recs"] if r["include"] and norm(r["title"]) not in template_titles]

    def anchor_before(*titles):
        texts = {ptext(el): el for el in children() if el.tag == qn("w:p")}
        return next(texts[t] for t in titles if t in texts)

    # md-only recs go before the IDE section; library sections before Reporting Management.
    groups = [([r for r in extra if not r.get("library")], ("Stimulate the adoption of SonarQube for IDE Usage", "Reporting Management", "Annexes")),
              ([r for r in extra if r.get("library")], ("Reporting Management", "Annexes"))]
    sec_h = units["Enforce Quality Gates"][1][0]
    sec_intro = units["Integrate with DevOps Platform(s)"][1][1]
    item_t, item_d = units["Add Branch Protection Rules to GitHub repositories"][1][:2]
    for group, anchors in groups:
        if not group:
            continue
        anchor = anchor_before(*anchors)
        last_sec = object()
        for r in group:
            if r["section"] != last_sec and r["section"]:
                anchor.addprevious(set_text(copy.deepcopy(sec_h), r["section"]))
                intro_text = recs_md_sections[r["section"]]["intro"]
                if intro_text:
                    p = copy.deepcopy(sec_intro)
                    set_text_with_links(p, intro_text, doc.part)
                    anchor.addprevious(p)
            last_sec = r["section"]
            anchor.addprevious(set_text(copy.deepcopy(item_t), r["title"]))
            p = copy.deepcopy(item_d)
            set_text_with_links(p, fmt(r["description_md"], ctx, version), doc.part)
            anchor.addprevious(p)

    # Drop the template's orange "fill me in" highlighting.
    for shd in doc.element.body.iter(qn("w:shd")):
        if (shd.get(qn("w:fill")) or "").lower() in TEMPLATE_HIGHLIGHTS:
            shd.getparent().remove(shd)
    keep_previous(ctx["out_docx"])
    doc.save(ctx["out_docx"])


def keep_previous(path):
    """Move an existing report to reports/previous/ so hand edits (pasted charts) survive a re-run."""
    path.parent.mkdir(exist_ok=True)
    if path.exists():
        prev = path.parent / "previous"
        prev.mkdir(exist_ok=True)
        stamp = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H%M")
        shutil.move(path, prev / f"{path.stem} ({stamp}){path.suffix}")


def between_inclusive(kids, start, end):
    return kids[kids.index(start): kids.index(end)]


# ---------------------------------------------------------------- report

def write_report(result, meta, ctx):
    lv = result["level"]
    sc = result["score"]
    out = [f"# {ctx['name']} – Sonar Maturity Assessment (Jev)", ""]
    files = ctx["files"]
    out += ["Inputs: " + "; ".join(f"{k}: {(v if isinstance(v, str) else ', '.join(v or [])) or 'none'}" for k, v in files.items())
            + (f"; {len(ctx['confirmed'])} confirmed facts" if ctx["confirmed"] else ""), ""]
    if meta:
        out += [f"Model: `{meta.get('model')}` · request_id: `{meta.get('request_id')}` · usage: {meta.get('usage')}", ""]
    out += [f"## Current level: **Level {lv} – {LEVELS[lv][0]}**", ""]
    if lv < 5:
        out += [f"Next step: Level {lv + 1} – {LEVELS[lv + 1][0]}", ""]
    out += ["| Criterion | Level | P(yes) |", "|---|---|---|"]
    for n, crits in LEVEL_CRITERIA.items():
        for cid, *_ in crits:
            out.append(f"| {cid} | {n} | {result['answers'][f'crit_{cid}']['noul']:.2f} |")
    probs = ", ".join(f"L{int(k) + 1}: {v:.2f}" for k, v in sorted(sc["probabilities"].items(), key=lambda kv: int(kv[0])))
    agree = "agrees" if result["score_level"] == lv else "**disagrees – review**"
    out += ["", f"Holistic Score cross-check: most likely Level {result['score_level']} ({agree}); "
            f"distribution {probs}; confidence {sc['confidence']:.2f}", ""]

    out += ["## Recommendations", "", "| Include | Target level | Recommendation | Maturity criteria | P(applies) | P(already done) | Excluded because | Review |", "|---|---|---|---|---|---|---|---|"]
    for r in sorted(result["recs"], key=lambda r: (not r["include"], r["level"])):
        out.append(f"| {'✅' if r['include'] else '—'} | {r['level']} | {r['title']} (`{r['key']}`) | {', '.join(r['criteria']) or '–'} | {r['p_applies']:.2f} | {r['p_done']:.2f} | {r['reason']} | {'⚠️' if r['review'] else ''} |")

    out += ["", "## Executive summary statements", "", "| Used | Statement | P(supported) |", "|---|---|---|"]
    version = result["choices"].get("version")
    for k, s in summary_candidates(ctx).items():
        out.append(f"| {'✅' if k in result['summary'] else '—'} | {fmt(s, ctx, version)} | {result['answers'][f'sum_{k}']['noul']:.2f} |")
    out += ["", "## Quality Gate findings", "", "| Kept | Finding | P(evidenced) |", "|---|---|---|"]
    for k, (f, _) in QG_FINDINGS.items():
        out.append(f"| {'✅' if k in result['qg'] else '—'} | {f} | {result['answers'][f'qg_{k}']['noul']:.2f} |")
    a = result["answers"]
    choices = {k: v for k, v in a.items() if k.startswith("choice_")}
    if choices:
        out += ["", "## Selected facts", "", "| Fact | Choice | Confidence |", "|---|---|---|"]
        for k, v in choices.items():
            out.append(f"| {k[len('choice_'):]} | {v.get('choice')} | {v.get('confidence', 0):.2f} |")
    out += ["", "## Current environment (as written)", ""]
    for heading, items in result["environment"].items():
        out.append(f"**{heading}**")
        out += [f"{'  ' * indent}- {text}" for text, indent in items]
        out.append("")
    env_q = {k: v for k, v in a.items() if k.startswith("env_")}
    if env_q:
        out += ["| Environment statement | P(supported) |", "|---|---|"]
        for heading, items in library.ENV_CANDIDATES.items():
            for key, text in items.items():
                if (q := f"env_{slug(heading)}_{key}") in a:
                    out.append(f"| {heading}: {text} | {a[q]['noul']:.2f} |")
    out += ["", f"Output document: `{ctx['out_docx']}`", ""]
    ctx["out_report"].write_text("\n".join(out))


def library_recs():
    out = []
    for sec in library.EXTRA_SECTIONS:
        for it in sec["items"]:
            out.append(dict(it, section=sec["section"], description_md=it["description"], library=True))
    return out


def load_recs():
    md_recs, md_sections = parse_recommendations_md(RECS_MD.read_text())
    for r in md_recs:
        r["description_md"] = r["description"]
    recs = md_recs + [dict(r, description_md=r["description"]) for r in TEMPLATE_ONLY_RECS] + library_recs()
    return recs, md_sections


def estimate_tokens(text):
    """Jev's token count, approximately: one per word, digit and punctuation mark.
    Within about 10% of the usage Jev reported for past runs (characters / 4 was up to 40%
    low on timestamped transcripts). It mostly errs high."""
    return len(re.findall(r"[A-Za-z]+|\d|[^\sA-Za-z\d]", text))


def plan(ctx, recs):
    """Jev state, questions, and estimated tokens: {"question": state + the longest
    question (TOKEN_BUDGET), "request": state + every question (REQUEST_BUDGET)}."""
    state = build_state(ctx)
    questions = build_questions(recs, state, ctx)
    q = [estimate_tokens(json.dumps(x.model_dump() if hasattr(x, "model_dump") else str(x), default=str)) for x in questions.values()]
    s = estimate_tokens(json.dumps(state))
    return state, questions, {"question": s + max(q), "request": s + sum(q)}


def over_budget(tokens):
    """Why the evidence is too large for one Jev call, or None when it fits."""
    if tokens["question"] > TOKEN_BUDGET:
        return (f"The evidence plus the longest question is ~{tokens['question']:,} tokens; Jev allows 32k "
                f"(budget {TOKEN_BUDGET:,}). Trim the transcript or notes, for example small talk.")
    if tokens["request"] > REQUEST_BUDGET:
        return (f"The evidence plus all questions is ~{tokens['request']:,} tokens; Jev allows 64k "
                f"(budget {REQUEST_BUDGET:,}). Trim the transcript or notes, for example small talk.")
    return None


def run(ctx, from_cache=False, log=lambda msg: print(msg, file=sys.stderr)):
    """Ask Jev (or reuse its cached answers), then write the docx and the markdown report."""
    recs, md_sections = load_recs()
    log(f"{ctx['name']}: notes={ctx['files']['notes'] or 'none'}, transcripts={ctx['files']['transcripts'] or 'none'}, "
        f"metrics={len(ctx['files']['metrics'])} CSVs, server info={ctx['files']['server_info'] or 'none'}")

    if from_cache:
        if not ctx["cache"].exists():
            raise AssessError(f"No cached Jev answers for {ctx['name']} yet: run once without --from-cache.")
        cached = json.loads(ctx["cache"].read_text())
        answers, meta = cached["answers"], cached.get("meta")
    else:
        state, questions, tokens = plan(ctx, recs)
        if problem := over_budget(tokens):
            raise AssessError(problem)
        log(f"Asking Jev {len(questions)} questions (~{tokens['question']:,} tokens for evidence + longest question, "
            f"~{tokens['request']:,} for the whole request)...")
        answers, meta = ask_jev(state, questions)
        ctx["cache"].write_text(json.dumps({"answers": answers, "meta": meta}, indent=2, default=str))

    result = decide(answers, recs, ctx)
    result["answers"] = answers
    build_docx(result, md_sections, ctx)
    write_report(result, meta, ctx)
    return result


def main():
    ap = argparse.ArgumentParser(description="SonarQube maturity assessment with Jev.")
    ap.add_argument("customer", nargs="?", help="customer folder name (or part of it) under customers/, or a path")
    ap.add_argument("--new", metavar="NAME", help="create a customer folder and exit")
    ap.add_argument("--assessor", help="with --new: the assessor's name (defaults to MATURITY_ASSESSOR in .env)")
    ap.add_argument("--list", action="store_true", help="list customer folders")
    ap.add_argument("--from-cache", action="store_true", help="reuse the customer's cached Jev answers")
    args = ap.parse_args()

    try:
        if args.new:
            new_customer(args.new, args.assessor)
            return
        if args.list or not args.customer:
            print("Customers: " + (", ".join(list_customers()) or "none yet (create one with --new NAME)"))
            return
        ctx = load_customer(find_customer(args.customer))
        run(ctx, from_cache=args.from_cache)
    except AssessError as e:
        sys.exit(str(e))
    print(ctx["out_report"].read_text())


if __name__ == "__main__":
    main()
